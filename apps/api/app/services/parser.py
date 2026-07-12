"""简历解析层（确定性文本抽取，可插拔工厂）。

M1：支持 PDF / Word(.docx) / 纯文本 / Markdown。
OCR(图片) 与 legacy .doc 在 MVP 范围之外（v0.3 决策），预留接口抛 NotImplementedError。

设计原则（对抗性审查结论）：
- 解析只产出"原始文本"，**不调用 LLM**——确定性、可单测、与后续"理解"环节解耦。
- 解析失败抛 ParseError，由 worker 捕获并标记 task failed，不影响整条链路的其他分支。
"""
from __future__ import annotations

import io
from abc import ABC, abstractmethod

from app.core.logging import get_logger

logger = get_logger("parser")


class ParseError(Exception):
    """解析失败（文件损坏/格式不支持/依赖缺失）。"""


class BaseParser(ABC):
    @abstractmethod
    async def parse(self, data: bytes) -> str: ...


class TextParser(BaseParser):
    """txt / md：直接解码，依次尝试常见编码。"""

    async def parse(self, data: bytes) -> str:
        for enc in ("utf-8", "gb18030", "latin-1"):
            try:
                return data.decode(enc)
            except UnicodeDecodeError:
                continue
        raise ParseError("无法解码文本文件（非 UTF-8/GB18030/Latin-1）")


class PdfParser(BaseParser):
    async def parse(self, data: bytes) -> str:
        try:
            from pdfminer.high_level import extract_text
        except ImportError as e:  # pragma: no cover
            raise ParseError("缺少 pdfminer.six 依赖，请联系运维") from e
        try:
            text = extract_text(io.BytesIO(data)) or ""
            if not text.strip():
                logger.warning("pdf_parse_empty", hint="可能是扫描件(图片 PDF)，需 OCR")
            return text
        except Exception as e:  # noqa: BLE001
            logger.warning("pdf_parse_failed", error=str(e))
            raise ParseError(f"PDF 解析失败: {e}") from e


class DocxParser(BaseParser):
    async def parse(self, data: bytes) -> str:
        try:
            from docx import Document
        except ImportError as e:  # pragma: no cover
            raise ParseError("缺少 python-docx 依赖，请联系运维") from e
        try:
            doc = Document(io.BytesIO(data))
            parts = [p.text for p in doc.paragraphs if p.text]
            for tbl in doc.tables:
                for row in tbl.rows:
                    cells = [c.text for c in row.cells if c.text]
                    if cells:
                        parts.append(" | ".join(cells))
            return "\n".join(parts)
        except Exception as e:  # noqa: BLE001
            logger.warning("docx_parse_failed", error=str(e))
            raise ParseError(f"DOCX 解析失败: {e}") from e


class LegacyDocParser(BaseParser):
    """legacy .doc (OLE) 解析复杂，MVP 不支持；明确报错而非静默失败。"""

    async def parse(self, data: bytes) -> str:
        raise NotImplementedError("老版 .doc 解析将在后续里程碑支持（建议转为 .docx）")


class OcrParser(BaseParser):
    """图片简历 OCR（v0.3 决策：MVP 范围外，预留云 OCR 接口）。"""

    async def parse(self, data: bytes) -> str:
        raise NotImplementedError("OCR 解析将在 MVP 之后接入（云 OCR 接口占位）")


_PARSERS: dict[str, type[BaseParser]] = {
    "text": TextParser,
    "pdf": PdfParser,
    "docx": DocxParser,
    "doc": LegacyDocParser,
    "ocr": OcrParser,
}


class ParserFactory:
    @staticmethod
    def get(file_type: str) -> BaseParser:
        cls = _PARSERS.get(file_type)
        if cls is None:
            raise ParseError(f"不支持的文件类型: {file_type}")
        return cls()

    @staticmethod
    def supported() -> list[str]:
        return list(_PARSERS.keys())
